import io
import json
import math
import sqlite3
from unittest.mock import Mock

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app import create_app
from knowledge import KnowledgeError, PgVectors, chunk_units, embeddings, extract_document
from test_app import setup, configure, add_repo, regular_user


def pdf_bytes(pages):
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        if text:
            stream = DecodedStreamObject()
            stream.set_data(f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def fake_model(url, **kwargs):
    response = Mock()
    if url.endswith('/embeddings'):
        values = []
        for i, text in enumerate(kwargs['json']['input']):
            vector = [1., 0., 0.] if any(word in text.lower() for word in ('leave', 'holiday', 'vacation')) else [0., 1., 0.]
            values.append({'index': i, 'embedding': vector})
        response.json.return_value = {'data': values}
    else:
        response.json.return_value = {'choices': [{'message': {'content': 'The allowance is 25 days [1].'}}]}
    return response


def enable_embeddings(client, headers, monkeypatch, model='embed-v1'):
    monkeypatch.setattr('knowledge.requests.post', fake_model)
    result = client.put('/api/admin/embedding-settings', headers=headers,
                        json={'base_url': 'http://localhost:11434/v1', 'model': model, 'dimensions': 3, 'api_key': 'embedding-test-secret'})
    assert result.status_code == 200, result.json


def upload_pdf(client, headers, rid, pages=None):
    return client.post('/api/admin/documents/upload', headers=headers, data={
        'repository_id': str(rid), 'file': (io.BytesIO(pdf_bytes(pages or ['Introduction', 'Annual leave allowance is 25 days.'])), 'policy.pdf')})


def test_pdf_pages_and_clickable_source_access(setup):
    app, client, headers = setup
    rid = client.post('/api/admin/repositories', headers=headers, json={'name': 'HR', 'groups': ['hr']}).json['id']
    result = upload_pdf(client, headers, rid)
    assert result.status_code == 201
    conn = sqlite3.connect(app.config['DATABASE'])
    rows = conn.execute('SELECT id,page,content FROM passages ORDER BY page').fetchall()
    conn.close()
    assert [row[1] for row in rows] == [1, 2]
    sid = rows[1][0]
    assert client.get('/sources/' + sid).status_code == 200
    source = client.get('/api/sources/' + sid).json
    assert source['label'] == 'Page 2'
    assert '25 days' in source['content']
    original = client.get('/api/sources/' + sid + '/file')
    assert original.mimetype == 'application/pdf' and original.data.startswith(b'%PDF')
    user, _ = regular_user(app)
    for path in ['/sources/', '/api/sources/']:
        assert user.get(path + sid).status_code == 404
    assert user.get('/api/sources/' + sid + '/file').status_code == 404
    assert app.test_client().get('/api/sources/' + sid + '/file').status_code == 401


def test_docx_sections_tables_and_no_invented_pages():
    doc = Document()
    doc.add_heading('Benefits', level=1)
    doc.add_paragraph('Employees can request flexible hours.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Annual leave'
    table.cell(0, 1).text = '25 days'
    raw = io.BytesIO()
    doc.save(raw)
    units, mime, warning = extract_document('benefits.docx', raw.getvalue())
    assert units[1]['label'] == 'Benefits / paragraph 2'
    assert '25 days' in units[2]['text'] and 'table 3' in units[2]['label']
    assert all(unit['page'] is None for unit in units)
    assert 'wordprocessingml' in mime


def test_scanned_and_invalid_documents():
    with pytest.raises(KnowledgeError, match='OCR'):
        extract_document('scan.pdf', pdf_bytes(['']))
    units, _, warning = extract_document('mixed.pdf', pdf_bytes(['', 'Searchable page']))
    assert units[0]['page'] == 2 and '1 page' in warning
    with pytest.raises(KnowledgeError):
        extract_document('fake.docx', b'not a zip')
    with pytest.raises(KnowledgeError):
        extract_document('malware.exe', b'hello')
    with pytest.raises(KnowledgeError):
        extract_document('bad.txt', b'\xff')


def test_chunking_never_crosses_pages():
    chunks = chunk_units([{'text': 'A' * 3000, 'page': 1, 'label': 'Page 1'},
                          {'text': 'B' * 2000, 'page': 2, 'label': 'Page 2'}])
    assert len(chunks) == 5
    assert all(len(c['content']) <= 1800 for c in chunks)
    assert all(('B' not in c['content']) for c in chunks if c['page'] == 1)


def test_semantic_search_citations_reindex_and_deletion(setup, monkeypatch):
    app, client, headers = setup
    configure(client, headers)
    enable_embeddings(client, headers, monkeypatch)
    rid, _ = add_repo(client, headers)
    uploaded = upload_pdf(client, headers, rid)
    assert uploaded.status_code == 201
    result = client.post('/api/chat', headers=headers, json={'message': 'How much vacation do I get?'})
    assert result.status_code == 200
    sources = result.json['sources']
    assert any(s['page'] == 2 and s['url'].startswith('/sources/') for s in sources)
    sid = next(s['source_id'] for s in sources if s['page'] == 2)
    did = uploaded.json['id']
    assert client.post(f'/api/admin/documents/{did}/reindex', headers=headers, json={}).status_code == 200
    assert client.get('/api/sources/' + sid).status_code == 200  # stable citations
    client.delete(f'/api/admin/documents/{did}', headers=headers)
    assert client.get('/api/sources/' + sid).status_code == 404
    assert client.get('/api/sources/' + sid + '/file').status_code == 404


def test_embedding_secret_masking_and_model_change(setup, monkeypatch):
    app, client, headers = setup
    enable_embeddings(client, headers, monkeypatch)
    rid, did = add_repo(client, headers)
    response = client.get('/api/admin/embedding-settings')
    assert response.json['has_api_key'] is True
    assert 'embedding-test-secret' not in response.get_data(as_text=True)
    assert 'embedding_key' not in client.get('/api/admin').json['settings']
    conn = sqlite3.connect(app.config['DATABASE'])
    saved = conn.execute('SELECT embedding_key FROM settings').fetchone()[0]
    conn.close()
    assert saved.startswith('gAAAA')
    enable_embeddings(client, headers, monkeypatch, model='embed-v2')
    docs = client.get('/api/admin').json['documents']
    assert docs[0]['status'] == 'needs_reindex'
    configure(client, headers)
    result = client.post('/api/chat', headers=headers, json={'message': 'Annual leave'})
    assert result.json['sources'] == []
    assert client.post(f'/api/admin/documents/{did}/reindex', headers=headers, json={}).status_code == 200
    assert client.get('/api/admin').json['documents'][0]['status'] == 'ready'


def test_failed_embedding_preserves_document_for_retry(setup, monkeypatch):
    _, client, headers = setup
    enable_embeddings(client, headers, monkeypatch)
    rid = client.post('/api/admin/repositories', headers=headers, json={'name': 'HR'}).json['id']
    bad = Mock()
    bad.json.return_value = {'data': [{'index': 0, 'embedding': [1, 2]}]}
    monkeypatch.setattr('knowledge.requests.post', Mock(return_value=bad))
    failed = upload_pdf(client, headers, rid)
    assert failed.status_code == 422
    docs = client.get('/api/admin').json['documents']
    assert docs[0]['status'] == 'failed' and 'dimensions' in docs[0]['error']
    monkeypatch.setattr('knowledge.requests.post', fake_model)
    assert client.post(f'/api/admin/documents/{failed.json["id"]}/reindex', headers=headers, json={}).status_code == 200


@pytest.mark.parametrize('vector', [[0, 0, 0], [1, float('nan'), 0], [1, float('inf'), 0], [1, 2], ['1', 0, 0]])
def test_bad_vectors_rejected_without_saving_settings(setup, monkeypatch, vector):
    _, client, headers = setup
    response = Mock()
    response.json.return_value = {'data': [{'index': 0, 'embedding': vector}]}
    monkeypatch.setattr('knowledge.requests.post', Mock(return_value=response))
    result = client.put('/api/admin/embedding-settings', headers=headers,
                        json={'base_url': 'http://local/v1', 'model': 'bad', 'dimensions': 3})
    assert result.status_code == 400
    assert client.get('/api/admin/embedding-settings').json['enabled'] is False


def test_semantic_acl_and_source_revocation(setup, monkeypatch):
    app, client, headers = setup
    configure(client, headers)
    enable_embeddings(client, headers, monkeypatch)
    rid, did = add_repo(client, headers, ['hr'])
    user, user_headers = regular_user(app, ['hr'])
    result = user.post('/api/chat', headers=user_headers, json={'message': 'Vacation allowance'})
    sid = result.json['sources'][0]['source_id']
    assert user.get('/api/sources/' + sid).status_code == 200
    conn = sqlite3.connect(app.config['DATABASE'])
    conn.execute("UPDATE users SET groups_json='[]' WHERE id='user-1'")
    conn.commit()
    conn.close()
    assert user.get('/api/sources/' + sid).status_code == 404
    assert user.post('/api/chat', headers=user_headers, json={'message': 'Vacation allowance'}).json['sources'] == []
    assert user.post('/api/admin/documents/upload', headers=user_headers, data={}).status_code == 403


def test_old_database_migration_and_restart_idempotence(setup):
    app, client, headers = setup
    rid, _ = add_repo(client, headers)
    conn = sqlite3.connect(app.config['DATABASE'])
    did = conn.execute("INSERT INTO documents(repo_id,title,content) VALUES(?, 'Old document', 'Legacy knowledge')", (rid,)).lastrowid
    conn.commit()
    conn.close()
    for _ in range(2):
        create_app(dict(app.config))
    conn = sqlite3.connect(app.config['DATABASE'])
    assert conn.execute('SELECT count(*) FROM passages WHERE doc_id=?', (did,)).fetchone()[0] == 1
    conn.close()


def test_pgvector_sql_is_scoped_and_parameterized(monkeypatch):
    store = PgVectors('unused', 'workspace-1')
    conn = Mock()
    conn.execute.return_value.fetchall.return_value = [('chunk-id', .9)]
    context = Mock()
    context.__enter__ = Mock(return_value=conn)
    context.__exit__ = Mock(return_value=False)
    monkeypatch.setattr(store, 'connect', lambda: context)
    assert store.search(3, 'model-signature', [7], [1, 0, 0]) == [('chunk-id', .9)]
    query, params = conn.execute.call_args.args
    assert 'repo_id=ANY(%s)' in query and 'namespace=%s' in query and 'signature=%s' in query
    assert params[1:4] == ('workspace-1', 'model-signature', [7])
    with pytest.raises(KnowledgeError):
        store.table('3; DROP TABLE users')
